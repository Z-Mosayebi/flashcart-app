"""Writes the binary test fixtures for web/lib/file-parsers.test.ts.

Run once (fixtures are committed):
    pip install openpyxl python-docx
    python scripts/make_upload_fixtures.py
"""

from pathlib import Path

from docx import Document
from openpyxl import Workbook

OUT = Path(__file__).resolve().parent.parent / "web" / "lib" / "__fixtures__"
OUT.mkdir(parents=True, exist_ok=True)

wb = Workbook()
ws = wb.active
ws.append(["Deutsch", "English"])
ws.append(["der Hund", "the dog"])
ws.append([None, None])  # a blank row the parser must drop
ws.append(["die Katze", "the cat"])
ws.append(["Anzahl", 3])  # a number cell the parser must stringify
wb.save(OUT / "vocab.xlsx")

doc = Document()
doc.add_heading("Dativ", level=1)
doc.add_paragraph("Ich helfe dem Mann & der Frau.")
doc.add_heading("Akkusativ", level=2)
doc.add_paragraph("Ich sehe den Mann.")
doc.save(OUT / "notes.docx")

print(f"Wrote fixtures to {OUT}")
